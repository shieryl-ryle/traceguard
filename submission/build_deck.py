"""Fill the AI Builder Cup proposal template with TraceGuard content (python-pptx, template-preserving).

Usage: python submission/build_deck.py <template.pptx> submission/TraceGuard-proposal.pptx submission/screenshots
Template: the organisers' Google Slides deck, downloaded via File > Download > PowerPoint.
Requires python-pptx. Logos and the C-3 follow-up shot in screenshots/ come from scratchpad deck_assets (video story.html symbols, live app).
"""

import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE, OUT, SHOTS = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
REPO_URL, APP_URL = "https://github.com/shieryl-ryle/traceguard", "https://traceguard-343683630064.asia-southeast1.run.app"
VIDEO_URL = "https://drive.google.com/file/d/15qQuek3NzyF3jAUqO0b61K3u1IBdCGOg/view"

FONT, FONT_MED, FONT_SEMI = "Google Sans Flex", "Google Sans Flex Medium", "Google Sans Flex SemiBold"
INK, MUTED, TINT, LINE = RGBColor(0x20, 0x21, 0x24), RGBColor(0x5F, 0x63, 0x68), RGBColor(0xF1, 0xF3, 0xF4), RGBColor(0xDA, 0xDC, 0xE0)
BLUE, GREEN, AMBER, RED, WHITE = (RGBColor(0x1A, 0x73, 0xE8), RGBColor(0x1E, 0x8E, 0x3E), RGBColor(0xF2, 0xA9, 0x00),
                                  RGBColor(0xD9, 0x30, 0x25), RGBColor(0xFF, 0xFF, 0xFF))

prs = Presentation(str(TEMPLATE))
S = list(prs.slides)


def set_runs(para, text):
    """Replace a paragraph's text, keeping its first run's formatting."""
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def title_only(slide, text=None):
    """Keep the template title box's first paragraph (its styling), drop the template prompts."""
    box = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    paras = box.text_frame.paragraphs
    if text:
        set_runs(paras[0], text)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    box.height = Inches(0.45)
    return box


def text(slide, x, y, w, h, items, size=12, color=INK, font=FONT, bold_first=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, bullets=False, space_after=4, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, item in enumerate(items if isinstance(items, list) else [items]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        segs = item if isinstance(item, list) else [(item, {})]
        if bullets:
            p.text = ""
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.18)))
            pPr.set("indent", str(-Inches(0.16)))
            bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
            pPr.append(bu)
        for seg, opt in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.size = Pt(opt.get("size", size))
            f.name = opt.get("font", FONT_SEMI if (bold_first and i == 0) or opt.get("bold") else font)
            f.color.rgb = opt.get("color", color)
            f.italic = opt.get("italic", False)
    return tb


def card(slide, x, y, w, h, fill=TINT, line=None, radius=0.08, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if name:
        shp.name = name
    return shp


def badge(slide, x, y, label, fill, d=0.34, size=12):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.name, r.font.color.rgb = Pt(size), FONT_SEMI, WHITE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return c


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(1.5)
    tail = ln.line._get_or_add_ln().makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.line._get_or_add_ln().append(tail)
    return ln


def picture(slide, path, x, y, w=None, h=None, frame=True):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    if frame:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(0.75)
    return pic


def cropped_h(slide, path, x, y, h, l=0.0, t=0.0, r=0.0, b=0.0, px=(2880, 1920)):
    """Like `cropped`, but sized to height h."""
    aspect = (px[0] * (1 - l - r)) / (px[1] * (1 - t - b))
    return cropped(slide, path, x, y, h * aspect, l, t, r, b, px)


def cropped(slide, path, x, y, w, l=0.0, t=0.0, r=0.0, b=0.0, px=(2880, 1920)):
    """Place an image cropped by fractions, sized to width w with the cropped aspect ratio kept."""
    pic = picture(slide, path, x, y, w=w)
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    aspect = (px[0] * (1 - l - r)) / (px[1] * (1 - t - b))
    pic.width, pic.height = Inches(w), Inches(w / aspect)
    return pic


# ---------- Slide 1: team details ----------
box = next(sh for sh in S[0].shapes if sh.has_text_frame and sh.text_frame.text.strip())
paras = box.text_frame.paragraphs
set_runs(paras[4], "Problem Statement: BFSI – Intelligent Risk, Fraud & Financial Experiences. Fraud analysts "
                   "reviewing one alert can't see the connected fund movements behind it, or prove them.")
for r in paras[4].runs:
    r.font.size = Pt(15)
for i in (2, 3):
    for r in paras[i].runs:
        r.font.size = Pt(16)

# ---------- Slide 2: brief ----------
s = S[1]
title_only(s)
text(s, 0.4, 1.3, 5.5, 2.3, [
    [("TraceGuard", {"bold": True, "size": 20, "color": BLUE})],
    "A graph-powered AI investigation assistant for fraud operations analysts at banks, e-wallets and fintechs.",
    "It traces suspicious fund movements across accounts and devices, then writes an evidence-linked case brief "
    "with Gemini on Vertex AI, where every claim is cited to a real transaction or rule result.",
], size=13, space_after=8)
pillars = [("1", BLUE, "Graph evidence", "Paths, time windows, sums and rule flags computed by code, never by the model."),
           ("2", GREEN, "Cited AI brief", "Observations, counterevidence and unknowns. Every ID is validated before display."),
           ("3", AMBER, "Honest by design", "\"Pattern warrants review\", never \"guilty\". It says what the data can't show.")]
for i, (n, colr, head, body) in enumerate(pillars):
    y = 1.35 + i * 1.12
    card(s, 6.15, y, 3.45, 0.98, name=f"pillar {n}")
    badge(s, 6.3, y + 0.13, n, colr)
    text(s, 6.75, y + 0.08, 2.75, 0.85, [[(head, {"bold": True, "size": 12})], body], size=10.5, space_after=2)
flow = [("In", "One alert", BLUE), ("Out", "A cited case brief", GREEN), ("Decision", "The analyst's", AMBER)]
for i, (k, v, colr) in enumerate(flow):
    x = 0.4 + i * 1.88
    card(s, x, 3.5, 1.58, 0.78, fill=WHITE, line=LINE, name=f"flow {i + 1}")
    text(s, x + 0.08, 3.55, 1.42, 0.7, [[(k.upper(), {"size": 9, "color": colr, "bold": True})], [(v, {"bold": True, "size": 12})]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=1)
    if i < len(flow) - 1:
        arrow(s, x + 1.6, 3.89, x + 1.86, 3.89)
text(s, 0.4, 4.6, 5.5, 0.7, [[("Investigation aid only: ", {"bold": True}),
                              ("it does not block transfers, declare guilt, file regulatory reports or give a risk score. "
                               "All data is synthetic.", {})]], size=10.5, color=MUTED)

# ---------- Slide 3: opportunities ----------
s = S[2]
title_only(s)
cols = [("How it differs", BLUE, [
            "Shows connections, not single transactions: a time-ordered money trail across accounts and devices.",
            "Separates duties: graph tools compute, Gemini explains, the analyst judges.",
            "Refuses to guess: unknowns past the data boundary are stated, not filled in."]),
        ("How it solves the problem", GREEN, [
            "Retrieves the connected evidence for an alert in one click.",
            "Turns it into a brief with observations, counterevidence, unknowns and next checks.",
            "Each citation highlights its record on the graph and timeline, so it can be checked in place."]),
        ("USP", AMBER, [
            "Every cited ID and figure is validated against tool output, and the checks are shown on screen (\"Verified before display\"). A failing brief is regenerated once, then blocked.",
            "Counterevidence is computed, not invented (e.g. senders not on the shared device).",
            "Careful language enforced in code: no accusations, probabilities or risk scores."])]
for i, (head, colr, pts) in enumerate(cols):
    x = 0.35 + i * 3.13
    card(s, x, 1.3, 2.98, 3.55, name=f"opportunity {i + 1}")
    badge(s, x + 0.15, 1.45, str(i + 1), colr)
    text(s, x + 0.58, 1.47, 2.3, 0.35, [[(head, {"bold": True, "size": 13})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, 1.95, 2.72, 2.85, pts, size=12, bullets=True, space_after=7)

# ---------- Slide 4: features ----------
s = S[3]
title_only(s)
feats = [("Case queue", "3 synthetic alert cases, clearly labelled \"synthetic data\""),
         ("Timeline", "Each case's transfers in time order; select rows to highlight them"),
         ("Relationship graph", "Accounts and devices; one arrow per transfer, so repeats never merge"),
         ("5 deterministic tools", "Account context, fund trace, shared devices, rapid forwarding, fan-in"),
         ("AI investigation", "Gemini on Vertex AI over retrieved evidence, with a visible tool log"),
         ("Evidence-linked brief", "Observations, counterevidence, unknowns, suggested checks"),
         ("Citations and export", "Click an evidence ID to highlight it on the graph; export the brief as Markdown"),
         ("Follow-up question", "Ask about the case; answered from the same evidence, or refused if it can't be"),
         ("Readable errors", "Timeout, quota or permission errors explained; graph stays usable")]
for i, (head, body) in enumerate(feats):
    col, row = i % 3, i // 3
    x, y = 0.35 + col * 3.13, 1.3 + row * 1.33
    card(s, x, y, 2.98, 1.18, name=f"feature {i + 1}")
    badge(s, x + 0.13, y + 0.14, str(i + 1), [BLUE, GREEN, AMBER][col], d=0.3, size=11)
    text(s, x + 0.52, y + 0.12, 2.38, 1.0, [[(head, {"bold": True, "size": 12})], body], size=10.5, space_after=2)

# ---------- Slide 5: process flow ----------
s = S[4]
title_only(s)
steps = [("Open alert", "Pick a case from the queue"),
         ("Inspect", "Graph and time-ordered timeline"),
         ("Run AI investigation", "Fixed workflow runs every graph tool"),
         ("Validate", "Citations, figures and wording checked"),
         ("Review brief", "Click citations to verify each claim"),
         ("Export", "Markdown brief for the case file")]
w, gap, y = 1.4, 0.18, 1.55
for i, (head, body) in enumerate(steps):
    x = 0.35 + i * (w + gap)
    fill = RGBColor(0xE8, 0xF0, 0xFE) if i in (2, 3) else TINT
    card(s, x, y, w, 1.6, fill=fill, name=f"step {i + 1}")
    badge(s, x + w / 2 - 0.17, y + 0.14, str(i + 1), BLUE if i not in (2, 3) else GREEN)
    text(s, x + 0.06, y + 0.55, w - 0.12, 1.0, [[(head, {"bold": True, "size": 11.5})], body], size=10,
         align=PP_ALIGN.CENTER, space_after=2)
    if i < len(steps) - 1:
        arrow(s, x + w + 0.02, y + 0.8, x + w + gap - 0.02, y + 0.8)
card(s, 0.35, 3.45, 9.3, 1.45, fill=WHITE, line=LINE, name="validation detail")
text(s, 0.55, 3.55, 9.0, 1.3, [
    [("Inside steps 3–4 ", {"bold": True, "size": 12}), ("(the trust boundary)", {"size": 12, "color": MUTED})],
    "Graph tools compute paths, windows, sums, rule flags and contrasts, so every number is deterministic.",
    "Gemini (gemini-3.5-flash, structured JSON) writes the brief from that evidence only, with no free-form queries and no generated code.",
    "The validator rejects any brief that cites an unknown ID, states a figure not in the tool output, or uses accusatory wording. "
    "It regenerates once, then shows a readable error instead of the brief.",
], size=10.5, bullets=False, space_after=3)

# ---------- Slide 6: wireframe / mock ----------
s = S[5]
title_only(s, "Wireframes/Mock diagrams of the proposed solution (optional): the live investigation page")
pic = picture(s, SHOTS / "slide2-c1-graph.png", 0.35, 1.3, h=3.95)
img_w, img_h = 3.95 * 1.5, 3.95
for n, (fx, fy) in enumerate([(0.16, 0.035), (0.215, 0.235), (0.26, 0.47), (0.595, 0.40), (0.20, 0.90)], start=1):
    badge(s, 0.35 + fx * img_w, 1.3 + fy * img_h, str(n), [BLUE, GREEN, AMBER, BLUE, GREEN][n - 1], d=0.3, size=11)
callouts = [("Case queue", "Synthetic cases, each labelled", BLUE),
            ("Case header", "Alert description and scenario", GREEN),
            ("Relationship graph", "Accounts (circles), devices (squares), one arrow per transfer", AMBER),
            ("AI investigation panel", "Brief, Verified panel and follow-up question, scrolling beside the graph", BLUE),
            ("Below the graph", "Selection panel for the clicked node, then the full timeline", GREEN)]
for i, (head, body, colr) in enumerate(callouts):
    y = 1.3 + i * 0.8
    badge(s, 6.45, y + 0.06, str(i + 1), colr, d=0.3, size=11)
    text(s, 6.85, y, 2.8, 0.75, [[(head, {"bold": True, "size": 12})], body], size=10.5, space_after=1)

# ---------- Slide 7: architecture ----------
s = S[6]
title_only(s)
card(s, 0.35, 1.4, 1.55, 1.2, fill=WHITE, line=LINE, name="browser")
text(s, 0.4, 1.5, 1.45, 1.0, [[("Browser", {"bold": True, "size": 12})], "Analyst or judge, public URL"], size=10,
     align=PP_ALIGN.CENTER)
run_box = card(s, 2.3, 1.3, 4.9, 3.95, fill=RGBColor(0xE8, 0xF0, 0xFE), name="cloud run")
picture(s, SHOTS / "logo-google-cloud.png", 2.45, 1.43, w=0.36, frame=False)
text(s, 2.88, 1.38, 4.2, 0.55, [[("Cloud Run · traceguard", {"bold": True, "size": 12.5})],
                                [("asia-southeast1 · max 2 instances · runtime SA traceguard-run", {"size": 10, "color": MUTED})]],
     space_after=0)
inner = [("Streamlit app (Python 3.12)", "Case queue · graph · timeline · brief · export", WHITE),
         ("Graph layer: NetworkX MultiDiGraph", "5 read-only tools · rule flags · contrasts (ADR-003)", WHITE),
         ("AI layer", "Fixed retrieval → Gemini JSON brief → citation, figure and wording validation", WHITE),
         ("Synthetic dataset", "Seeded, bundled in the image · no database", WHITE)]
for i, (head, body, fill) in enumerate(inner):
    y = 2.05 + i * 0.78
    card(s, 2.5, y, 4.5, 0.66, fill=fill, line=LINE, name=f"layer {i + 1}")
    text(s, 2.6, y + 0.05, 4.3, 0.58, [[(head, {"bold": True, "size": 11})], body], size=9.5, space_after=0)
card(s, 7.6, 1.3, 2.05, 1.75, fill=WHITE, line=LINE, name="vertex ai")
picture(s, SHOTS / "logo-gemini.png", 8.27, 1.38, w=0.32, frame=False)
picture(s, SHOTS / "logo-vertex-ai.png", 8.67, 1.38, w=0.32, frame=False)
text(s, 7.65, 1.74, 1.95, 1.3, [[("Gemini on Vertex AI", {"bold": True, "size": 12})],
                                "gemini-3.5-flash", "asia-southeast1", "google-genai SDK", "structured JSON output"],
     size=10, align=PP_ALIGN.CENTER, space_after=0)
arrow(s, 1.9, 2.0, 2.3, 2.0)
arrow(s, 7.2, 2.15, 7.6, 2.15)
text(s, 7.15, 2.25, 0.5, 0.3, [[("ADC", {"size": 9, "color": MUTED})]], align=PP_ALIGN.CENTER)
text(s, 7.6, 3.15, 2.05, 2.1, [
    [("Security and cost", {"bold": True, "size": 11.5})],
    "No API keys: Application Default Credentials only",
    "Runtime SA: aiplatform.user only",
    "Build SA: run.builder only · no Editor role in the project",
    "Max instances plus per-call caps (30 s, 4 tool rounds)",
], size=9.5, bullets=False, space_after=2)

# ---------- Slide 8: technologies ----------
s = S[7]
title_only(s)
rows = [("Layer", "Technology", "Why"),
        ("AI", "Gemini 3.5 Flash on Vertex AI · google-genai 2.31", "Structured JSON briefs; LOW thinking: a brief in 3–9 s (measured)"),
        ("Hosting", "Cloud Run (source deploy, Dockerfile)", "One public service; scales to zero; max 2 instances"),
        ("Build", "Cloud Build · Artifact Registry", "Least-privilege build service account"),
        ("Graph", "NetworkX 3.7 (MultiDiGraph)", "Repeated transfers stay separate; time-ordered traces"),
        ("UI", "Streamlit 1.66 · streamlit-agraph", "Interactive graph; click a node to see its evidence"),
        ("Validation", "Pydantic 2 schema · custom validators", "Citation, figure and wording checks before display"),
        ("Quality", "pytest (168 tests) · mypy strict", "Fixture outcomes, boundaries, errors, honesty rules")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.35), Inches(1.3), Inches(9.3), Inches(3.9)).table
for ci, wdt in enumerate((1.25, 3.85, 4.2)):
    tbl.columns[ci].width = Inches(wdt)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if ri == 0 else (TINT if ri % 2 == 0 else WHITE)
        tf = cell.text_frame
        tf.paragraphs[0].text = ""
        r = tf.paragraphs[0].add_run()
        r.text = val
        r.font.size = Pt(11 if ri == 0 else 10.5)
        r.font.name = FONT_SEMI if ri == 0 or ci == 0 else FONT
        r.font.color.rgb = WHITE if ri == 0 else INK
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
text(s, 0.35, 5.22, 9.3, 0.3, [[("Two of the theme's suggested Google Cloud technologies (Gemini on Vertex AI, Cloud Run); "
                                 "BigQuery is the scale-up path.", {"size": 9.5, "color": MUTED})]])

# ---------- Slide 9: cost ----------
s = S[8]
title_only(s)
text(s, 0.4, 1.3, 9.2, 0.6, [[("We have not published a cost figure: we would rather show the drivers and the caps than guess a number.",
                               {"size": 12})]])
drivers = [("Cloud Run", "Pay per request time; scales to zero when idle; capped at 2 instances.", BLUE),
           ("Gemini 3.5 Flash", "One call per investigation (two at most, with one regeneration). Evidence payload capped at 60,000 characters.", GREEN),
           ("Build and storage", "Cloud Build per deploy; one container image in Artifact Registry.", AMBER)]
for i, (head, body, colr) in enumerate(drivers):
    x = 0.35 + i * 3.13
    card(s, x, 2.0, 2.98, 1.75, name=f"cost driver {i + 1}")
    badge(s, x + 0.15, 2.15, str(i + 1), colr)
    text(s, x + 0.58, 2.15, 2.3, 1.5, [[(head, {"bold": True, "size": 12.5})], body], size=10.5, space_after=3)
text(s, 0.4, 3.95, 9.2, 1.3, [
    [("Controls in place: ", {"bold": True}), ("max instances 2; 30 s timeout per call; at most 4 tool rounds and 1 regeneration; "
                                               "budget alerts at 50/90/100%.", {})],
    [("At bank scale: ", {"bold": True}), ("cost would be driven by Gemini calls per analyst case and graph storage "
                                           "(BigQuery or Spanner Graph). Estimate with the bank's case volumes before rollout.", {})],
], size=11, space_after=6)

# ---------- Slide 10: snapshots ----------
s = S[9]
title_only(s)
H = 2.7
shots = [
    (SHOTS / "slide3-brief.png", dict(l=0.255, t=0.39, r=0.40, b=0.0), (2880, 1920),
     [("Citation click: ", {"bold": True}), ("records cited by RULE-FWD-01 highlighted on the graph (amber).", {})]),
    (SHOTS / "slide10-verified-panel.png", dict(l=0.0, t=0.0, r=0.05, b=0.0), (890, 600),
     [("Verified before display: ", {"bold": True}), ("citations, figures and wording checked, shown beside the brief.", {})]),
    (SHOTS / "slide10-follow-up-refusal.png", dict(l=0.0, t=0.0, r=0.0, b=0.0), (948, 874),
     [("Follow-up on C-3: ", {"bold": True}), ("\"Who owns A-22?\" is refused: the case data can't answer it.", {})]),
]
x = 0.35
for path, crop, px, caption in shots:
    pic = cropped_h(s, path, x, 1.3, H, px=px, **crop)
    w = pic.width / 914400
    text(s, x, 1.3 + H + 0.08, w, 0.75, [caption], size=10, color=MUTED)
    x += w + 0.15

# ---------- Slide 11: performance ----------
s = S[10]
title_only(s)
stats = [("3 / 3", "cases reach the expected\nconclusion on the live URL"), ("0", "invented citations shown\n(blocked by design)"),
         ("3–9 s", "live brief (LOW thinking),\nvalid on the first attempt"), ("11 / 11", "graph-tool checks match\nexpected outcomes")]
for i, (big, small) in enumerate(stats):
    x = 0.35 + i * 2.35
    card(s, x, 1.3, 2.2, 1.25, name=f"stat {i + 1}")
    text(s, x + 0.1, 1.33, 2.0, 0.6, [[(big, {"bold": True, "size": 24, "color": BLUE})]], align=PP_ALIGN.CENTER)
    text(s, x + 0.1, 1.92, 2.0, 0.6, small.split("\n"), size=9.5, color=MUTED, align=PP_ALIGN.CENTER, space_after=0)
rows = [("Case", "Pattern", "Expected", "Live result (Gemini 3.5 Flash, real graph evidence)"),
        ("C-1", "Convergence & forwarding", "Pattern warrants review", "Both rules flag; brief cites 4 senders, ratio 0.91, 30 min; counterevidence names A-02/A-04"),
        ("C-2", "Benign shared device", "Insufficient evidence", "No rule matches; device D-04 shown as an observation, not as proof"),
        ("C-3", "Incomplete evidence", "Unknown past the boundary", "Next recipient after A-22 stated as unknown; nothing guessed")]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.35), Inches(2.75), Inches(9.3), Inches(1.75)).table
for ci, wdt in enumerate((0.6, 1.9, 1.95, 4.85)):
    tbl.columns[ci].width = Inches(wdt)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE if ri == 0 else (TINT if ri % 2 == 0 else WHITE)
        tf = cell.text_frame
        tf.paragraphs[0].text = ""
        r = tf.paragraphs[0].add_run()
        r.text = val
        r.font.size = Pt(10 if ri == 0 else 9.5)
        r.font.name = FONT_SEMI if ri == 0 or ci == 0 else FONT
        r.font.color.rgb = WHITE if ri == 0 else INK
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
text(s, 0.35, 4.62, 9.3, 0.7, [
    [("Latency fix, measured: ", {"bold": True}), ("default thinking 20–28 s per brief (one 504 at the 30 s cap) → LOW thinking 3–9 s, "
                                                   "all valid. 168 automated tests; mypy strict. ", {}),
     ("No production accuracy or time-saved claims: the data is synthetic.", {"italic": True})],
], size=9.5, color=MUTED)

# ---------- Slide 12: future ----------
s = S[11]
title_only(s)
future = [("Scale", "Spanner Graph or BigQuery for bank-scale transaction graphs, replacing in-memory NetworkX.", BLUE),
          ("Access and audit", "Analyst sign-in, case-level access controls, and an audit log of every brief and citation click.", GREEN),
          ("Validation", "Back-test on the bank's labelled historical alerts, under its governance. Track analyst agreement before any accuracy claim.", AMBER),
          ("More patterns", "Cycle detection, CSV import of the bank's own alerts, and cross-case links between accounts (cut from this build).", BLUE)]
for i, (head, body, colr) in enumerate(future):
    col, row = i % 2, i // 2
    x, y = 0.35 + col * 4.7, 1.3 + row * 1.6
    card(s, x, y, 4.55, 1.45, name=f"future {i + 1}")
    badge(s, x + 0.15, y + 0.15, str(i + 1), colr)
    text(s, x + 0.6, y + 0.13, 3.85, 1.25, [[(head, {"bold": True, "size": 12.5})], body], size=10.5, space_after=3)
text(s, 0.35, 4.6, 9.3, 0.6, [[("Human-led: ", {"bold": True}),
                               ("the analyst always decides. TraceGuard prepares the evidence and the brief.", {})]], size=11, color=MUTED)

# ---------- Slide 13: links ----------
s = S[12]
title_only(s, "Provide links to your:")
links = [("GitHub Public Repository", REPO_URL, BLUE), ("Demo Video Link (3–4 Minutes)", VIDEO_URL, GREEN),
         ("Final Product Link", APP_URL, AMBER)]
for i, (head, url, colr) in enumerate(links):
    y = 1.35 + i * 1.12
    card(s, 0.35, y, 9.3, 0.95, name=f"link {i + 1}")
    badge(s, 0.52, y + 0.3, str(i + 1), colr)
    tb = text(s, 1.0, y + 0.1, 8.5, 0.8, [[(head, {"bold": True, "size": 13})], [(url, {"size": 12, "color": BLUE})]], space_after=2)
    if url.startswith("http"):
        tb.text_frame.paragraphs[1].runs[0].hyperlink.address = url
text(s, 0.35, 4.75, 9.3, 0.5, [[("Live app runs C-1, C-2 and C-3 end to end, including a real Gemini call. No sign-in needed. "
                                 "Video: 3:56, follows a fictional analyst through the three synthetic cases.",
                                 {"size": 10.5, "color": MUTED})]])

# ---------- Slide 15: thank you (close on the product's promise) ----------
s = S[14]
GREY = RGBColor(0xBD, 0xC1, 0xC6)
text(s, 6.75, 2.3, 3.0, 2.6, [
    [("Trace", {"bold": True, "size": 28, "color": WHITE}), ("Guard", {"bold": True, "size": 28, "color": RGBColor(0x34, 0xA8, 0x53)})],
    [("Follow the connections.", {"size": 15, "color": WHITE})],
    [("Keep the judgment human.", {"size": 15, "color": WHITE})],
    [("Team Chilli · AI Builder Cup 2026", {"size": 11, "color": GREY})],
    [(APP_URL.replace("https://", ""), {"size": 10, "color": GREY})],
    [(REPO_URL.replace("https://", ""), {"size": 10, "color": GREY})],
], space_after=6)

# ---------- Slide 14: limitations (template slide without a title) ----------
s = S[13]
title_box = copy.deepcopy(next(sh for sh in S[1].shapes if sh.has_text_frame and sh.text_frame.text.strip())._element)
s.shapes._spTree.append(title_box)
t = s.shapes[-1]
set_runs(t.text_frame.paragraphs[0], "Limitations and responsible AI")
for p in t.text_frame.paragraphs[1:]:
    p._p.getparent().remove(p._p)
limits = [("Limitations", RED, [
              "Synthetic data only: 35 accounts, 102 transfers, one currency (PHP).",
              "No measured accuracy or time saved yet.",
              "Fixed workflow: the model does not choose tools (stated plainly, not as autonomy).",
              "No authentication in the prototype.",
              "Model wording varies between runs; the validator rejects bad output rather than preventing it."]),
          ("Safeguards built in", GREEN, [
              "Synthetic data only; no employer or customer data (BR-001).",
              "Every cited ID must exist in retrieved evidence (BR-002).",
              "Numbers and flags from code, not the model (BR-003).",
              "No accusations, probabilities or risk scores (BR-004).",
              "Unknowns stated; similar amounts are never called the same money (BR-005).",
              "Live vs cached output always labelled (BR-006)."])]
for i, (head, colr, pts) in enumerate(limits):
    x = 0.35 + i * 4.7
    card(s, x, 1.3, 4.55, 3.7, name=f"limits {i + 1}")
    badge(s, x + 0.15, 1.45, "!" if i == 0 else "✓", colr)
    text(s, x + 0.6, 1.47, 3.8, 0.35, [[(head, {"bold": True, "size": 13})]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, 1.95, 4.25, 3.0, pts, size=12, bullets=True, space_after=6)

prs.save(str(OUT))
print("saved", OUT)
